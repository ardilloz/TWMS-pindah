"""Criterion: Report PDF/Excel exports (backend refactored into helpers) unchanged.

Fetches /api/report/export.pdf and /api/report/export.xlsx with date filters covering the
seeded ZTEST transactions, and asserts the expected content/sheets are present. Also checks
401 without login. Does not create or mutate any data (export is read-only over seeded data).
"""
import io

import httpx
from openpyxl import load_workbook
from pypdf import PdfReader

from .conftest import API_URL


def test_report_export_pdf(admin_client):
    resp = admin_client.get(
        "/report/export.pdf",
        params={"date_from": "2026-10-01", "date_to": "2026-10-03"},
    )
    assert resp.status_code == 200, f"export.pdf failed: {resp.status_code} {resp.text[:300]}"
    assert resp.headers["content-type"] == "application/pdf"
    body = resp.content
    assert body[:4] == b"%PDF", "response is not a PDF"
    reader = PdfReader(io.BytesIO(body))
    assert len(reader.pages) >= 1
    text = "\n".join((p.extract_text() or "") for p in reader.pages)
    assert "Laporan Stok Gudang" in text, text[:500]
    assert "TWMS" in text, text[:500]
    assert "ZTEST Induk" in text, text[:500]
    assert "Subtotal" in text, text[:500]
    assert "Kiriman Supplier" in text, text[:500]

    with httpx.Client(base_url=API_URL, timeout=30.0) as anon:
        unauth = anon.get("/report/export.pdf")
        assert unauth.status_code == 401, f"no-login should 401: {unauth.status_code}"


def test_report_export_xlsx(admin_client):
    resp = admin_client.get(
        "/report/export.xlsx",
        params={"date_from": "2026-10-01", "date_to": "2026-10-03"},
    )
    assert resp.status_code == 200, f"export.xlsx failed: {resp.status_code} {resp.text[:300]}"
    assert "spreadsheet" in resp.headers["content-type"] or "excel" in resp.headers["content-type"]
    wb = load_workbook(io.BytesIO(resp.content))
    assert "Stok & Penggunaan" in wb.sheetnames, wb.sheetnames
    assert "Kiriman Supplier" in wb.sheetnames, wb.sheetnames

    ws = wb["Stok & Penggunaan"]
    header_cell = ws["A5"]  # row 1-4 are title/period/notes; row 5 is the column header row
    assert header_cell.value == "Gudang", header_cell.value
    # brown bold header styling expected
    assert header_cell.font.bold is True, "header should be bold"
    fill_rgb = (header_cell.fill.fgColor.rgb or "").upper()
    assert fill_rgb not in ("", "00000000", "FFFFFFFF"), f"expected a brown fill, got {fill_rgb}"

    with httpx.Client(base_url=API_URL, timeout=30.0) as anon:
        unauth = anon.get("/report/export.xlsx")
        assert unauth.status_code == 401, f"no-login should 401: {unauth.status_code}"
