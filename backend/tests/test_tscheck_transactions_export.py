"""Criterion: Riwayat Excel export (moved to lib/receipt.py refactor) unchanged.

Fetches /api/transactions/export and asserts the two expected sheets exist. Read-only over
seeded data; also checks 401 without login.
"""
import io

import httpx
from openpyxl import load_workbook

from .conftest import API_URL


def test_transactions_export_xlsx(admin_client):
    resp = admin_client.get("/transactions/export")
    assert resp.status_code == 200, f"export failed: {resp.status_code} {resp.text[:300]}"
    assert "spreadsheet" in resp.headers["content-type"] or "excel" in resp.headers["content-type"]
    wb = load_workbook(io.BytesIO(resp.content))
    assert "Riwayat Transaksi" in wb.sheetnames, wb.sheetnames
    assert "Stok Per Gudang" in wb.sheetnames, wb.sheetnames

    with httpx.Client(base_url=API_URL, timeout=30.0) as anon:
        unauth = anon.get("/transactions/export")
        assert unauth.status_code == 401, f"no-login should 401: {unauth.status_code}"
