"""Tiny Excel helpers (openpyxl) for the Internal module import/export."""

from io import BytesIO

from openpyxl import Workbook, load_workbook


def build_xlsx(headers: list[str], rows: list[list], sheet_name: str = "Data") -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name[:31]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = cell.font.copy(bold=True)
    for r in rows:
        ws.append(r)
    for i, h in enumerate(headers, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = max(12, min(40, len(str(h)) + 4))
    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue()


def read_xlsx(data: bytes) -> list[dict]:
    wb = load_workbook(BytesIO(data), data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(h).strip().lower().replace(" ", "_") if h is not None else f"col{i}" for i, h in enumerate(rows[0])]
    out: list[dict] = []
    for r in rows[1:]:
        if all(c is None or str(c).strip() == "" for c in r):
            continue
        out.append({headers[i]: (r[i] if i < len(r) else None) for i in range(len(headers))})
    return out
