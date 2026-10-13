"""Shared openpyxl helpers for the Excel exports."""

import io

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="3D2411")
BOLD = Font(bold=True)
XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def style_sheet(sheet: Worksheet, header_row: int = 1, min_width: int = 0, max_width: int = 45) -> None:
    """Brown header row + auto column widths computed from header_row downwards."""
    for c in sheet[header_row]:
        c.font, c.fill = HEADER_FONT, HEADER_FILL
    for col in sheet.iter_cols(min_row=header_row):
        width = max(len(str(c.value or "")) for c in col) + 2
        sheet.column_dimensions[col[0].column_letter].width = min(max(width, min_width), max_width)


def bold_last_row(sheet: Worksheet) -> None:
    for c in sheet[sheet.max_row]:
        c.font = BOLD


def workbook_bytes(wb: Workbook) -> io.BytesIO:
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
