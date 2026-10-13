import io
from datetime import datetime
from typing import Optional
from zoneinfo import ZoneInfo

from fastapi import APIRouter, Depends
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from lib.auth import allowed_warehouses, require_any
from lib.dates import fmt_ddmmyy
from lib.report import ITEMS, build_report, iter_report_rows, supplier_grand_total
from lib.xlsx import XLSX_MEDIA_TYPE, bold_last_row, style_sheet, workbook_bytes
from models.schemas import ReportOut

router = APIRouter(prefix="/report", tags=["report"])

ITEM_LABEL = {"glangse": "Glangse", "tali": "Tali Putih", "kertas": "Kertas Coklat"}
ITEM_UNIT = {"glangse": "gulung", "tali": "kg", "kertas": "roll"}
COLS = [("masuk", "Masuk Supplier"), ("transfer_in", "Transfer Masuk"), ("transfer_out", "Transfer Keluar"), ("pribadi", "Pemakaian"), ("stok", "Stok Akhir")]
BROWN = colors.HexColor("#3D2411")
LINE = colors.HexColor("#C9B9A3")
SUBTOTAL_BG = "#FBEFD9"
TOTAL_BG = "#EADCC6"


# ---------- formatting ----------
def nf(x: float, dec: int = 2) -> str:
    s = f"{x:,.{dec}f}".replace(",", "_").replace(".", ",").replace("_", ".")
    if "," in s:
        s = s.rstrip("0").rstrip(",")
    return s


def fmt(item: str, v: float) -> str:
    if item == "glangse":
        return f"{nf(v / 5)} gulung ({nf(v)} koli)"
    return f"{nf(v)} {ITEM_UNIT[item]}"


def to_display_unit(item: str, v: float) -> float:
    """Glangse is reported in gulung, other items in their base unit."""
    return round(v / 5, 4) if item == "glangse" else v


def period_label(date_from: Optional[str], date_to: Optional[str]) -> str:
    if not date_from and not date_to:
        return "Semua periode"
    return f"{fmt_ddmmyy(date_from) if date_from else 'awal'} s/d {fmt_ddmmyy(date_to) if date_to else 'sekarang'}"


def _filename(ext: str, date_from: Optional[str], date_to: Optional[str]) -> str:
    return f"twms-laporan-{date_from or 'awal'}-sd-{date_to or 'akhir'}.{ext}"


@router.get("", response_model=ReportOut, dependencies=[Depends(require_any("stok"))])
async def get_report(date_from: Optional[str] = None, date_to: Optional[str] = None, user: dict = Depends(require_any("stok"))):
    return ReportOut(**(await build_report(date_from, date_to, allowed_warehouses(user))))


# ---------- Excel ----------
def _xlsx_warehouse_sheet(ws: Worksheet, r: dict, period: str, printed_by: str) -> None:
    ws.title = "Stok & Penggunaan"
    ws.append(["Laporan Stok Gudang — TWMS"])
    ws.append([f"Periode: {period}", f"Dicetak oleh: {printed_by}"])
    ws.append(["Catatan: Glangse dalam gulung (1 gulung = 5 koli); kolom terakhir = stok akhir glangse dalam koli"])
    ws.append([])
    ws.append(["Gudang", "Jenis", "Gudang Induk", "Barang", "Satuan"] + [c[1] for c in COLS] + ["Stok Akhir (koli)"])
    header_row = ws.max_row
    for row, kind in iter_report_rows(r):
        for it in ITEMS:
            vals = [to_display_unit(it, row[k][it]) for k, _ in COLS]
            lembar = row["stok"][it] if it == "glangse" else None
            ws.append([row["warehouse_name"], kind, row.get("parent_name") or "", ITEM_LABEL[it], ITEM_UNIT[it], *vals, lembar])
            if kind in ("Subtotal", "Total"):
                bold_last_row(ws)
    style_sheet(ws, header_row, min_width=10, max_width=40)
    ws["A1"].font = Font(bold=True, size=14)


def _supplier_cells(t: dict) -> list:
    return [to_display_unit("glangse", t["glangse"]), t["glangse"], t["tali"], t["kertas"]]


def _xlsx_supplier_sheet(ws: Worksheet, r: dict) -> None:
    ws.append(["Supplier", "Gudang", "Glangse (gulung)", "Glangse (koli)", "Tali Putih (kg)", "Kertas Coklat (roll)"])
    for s in r["suppliers"]:
        for pw in s["per_warehouse"]:
            ws.append([s["supplier_name"], pw["warehouse_name"], *_supplier_cells(pw["totals"])])
        ws.append([s["supplier_name"], "TOTAL ke seluruh gudang", *_supplier_cells(s["total"])])
        bold_last_row(ws)
    ws.append(["SEMUA SUPPLIER", "TOTAL ke seluruh gudang", *_supplier_cells(supplier_grand_total(r))])
    bold_last_row(ws)
    style_sheet(ws, 1, min_width=10, max_width=40)


@router.get("/export.xlsx")
async def export_xlsx(date_from: Optional[str] = None, date_to: Optional[str] = None, user: dict = Depends(require_any("stok"))):
    r = await build_report(date_from, date_to, allowed_warehouses(user))
    wb = Workbook()
    _xlsx_warehouse_sheet(wb.active, r, period_label(date_from, date_to), user["name"])
    _xlsx_supplier_sheet(wb.create_sheet("Kiriman Supplier"), r)
    return StreamingResponse(
        workbook_bytes(wb),
        media_type=XLSX_MEDIA_TYPE,
        headers={"Content-Disposition": f'attachment; filename="{_filename("xlsx", date_from, date_to)}"'},
    )


# ---------- PDF ----------
def _pdf_styles() -> dict[str, ParagraphStyle]:
    ss = getSampleStyleSheet()
    cell = ParagraphStyle("cell", parent=ss["Normal"], fontSize=8, leading=10)
    return {
        "h1": ParagraphStyle("h1", parent=ss["Title"], alignment=0, fontSize=18, textColor=BROWN, spaceAfter=2),
        "h2": ParagraphStyle("h2", parent=ss["Heading2"], textColor=BROWN, fontSize=13, spaceBefore=10, spaceAfter=6),
        "small": ParagraphStyle("small", parent=ss["Normal"], fontSize=9, textColor=colors.HexColor("#57534E")),
        "cell": cell,
        "cellb": ParagraphStyle("cellb", parent=cell, fontName="Helvetica-Bold"),
    }


def _base_table_style(extra: list) -> TableStyle:
    return TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), BROWN),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, 0), 8.5),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("BOX", (0, 0), (-1, -1), 0.6, LINE),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        *extra,
    ])


def _pdf_row_label(row: dict, kind: str) -> str:
    if kind == "Ranting":
        return f"&nbsp;&nbsp;&nbsp;» {row['warehouse_name']}<br/><font size=7 color='#78716C'>Ranting</font>"
    if kind == "Subtotal":
        return f"Subtotal {row['warehouse_name']}"
    return row["warehouse_name"]


def _pdf_warehouse_table(r: dict, st: dict) -> Table:
    data: list = [["Gudang", "Barang"] + [c[1] for c in COLS]]
    extra: list = []
    for row, kind in iter_report_rows(r):
        style = st["cellb"] if kind in ("Subtotal", "Total") else st["cell"]
        start = len(data)
        for j, it in enumerate(ITEMS):
            label = _pdf_row_label(row, kind) if j == 0 else ""
            data.append([Paragraph(label, style), Paragraph(ITEM_LABEL[it], style)] + [Paragraph(fmt(it, row[k][it]), style) for k, _ in COLS])
        extra.append(("LINEABOVE", (0, start), (-1, start), 0.6, LINE))
        bg = {"Subtotal": SUBTOTAL_BG, "Total": TOTAL_BG}.get(kind)
        if bg:
            extra.append(("BACKGROUND", (0, start), (-1, len(data) - 1), colors.HexColor(bg)))
    if not r["rows"]:
        data.insert(1, [Paragraph("Belum ada gudang.", st["cell"]), "", "", "", "", "", ""])
    t = Table(data, colWidths=[62 * mm, 26 * mm] + [37.8 * mm] * 5, repeatRows=1)
    t.setStyle(_base_table_style(extra))
    return t


def _pdf_supplier_table(r: dict, st: dict) -> Table:
    cell, cellb = st["cell"], st["cellb"]
    data: list = [["Supplier", "Gudang Tujuan", "Glangse", "Tali Putih", "Kertas Coklat"]]
    extra: list = []
    for s in r["suppliers"]:
        extra.append(("LINEABOVE", (0, len(data)), (-1, len(data)), 0.6, LINE))
        for j, pw in enumerate(s["per_warehouse"]):
            name = s["supplier_name"] if j == 0 else ""
            data.append([Paragraph(name, cell), Paragraph(pw["warehouse_name"], cell)] + [Paragraph(fmt(it, pw["totals"][it]), cell) for it in ITEMS])
        data.append([Paragraph("", cell), Paragraph("Total ke seluruh gudang", cellb)] + [Paragraph(fmt(it, s["total"][it]), cellb) for it in ITEMS])
        extra.append(("BACKGROUND", (0, len(data) - 1), (-1, len(data) - 1), colors.HexColor(SUBTOTAL_BG)))
    grand = supplier_grand_total(r)
    data.append([Paragraph("Semua Supplier", cellb), Paragraph("Total ke seluruh gudang", cellb)] + [Paragraph(fmt(it, grand[it]), cellb) for it in ITEMS])
    extra.append(("BACKGROUND", (0, len(data) - 1), (-1, len(data) - 1), colors.HexColor(TOTAL_BG)))
    t = Table(data, colWidths=[70 * mm, 70 * mm, 47 * mm, 42 * mm, 44 * mm], repeatRows=1)
    t.setStyle(_base_table_style(extra))
    return t


def _pdf_footer(canvas, doc) -> None:
    canvas.saveState()
    canvas.setFont("Helvetica", 7.5)
    canvas.setFillColor(colors.HexColor("#78716C"))
    canvas.drawString(12 * mm, 7 * mm, "TWMS — Tobacco Warehouse Management System")
    canvas.drawRightString(landscape(A4)[0] - 12 * mm, 7 * mm, f"Halaman {doc.page}")
    canvas.restoreState()


def build_report_pdf(r: dict, period: str, printed_by: str) -> io.BytesIO:
    st = _pdf_styles()
    now = datetime.now(ZoneInfo("Asia/Jakarta")).strftime("%d/%m/%Y %H:%M")
    sep = "&nbsp;&nbsp;|&nbsp;&nbsp;"
    story = [
        Paragraph("Laporan Stok Gudang — TWMS", st["h1"]),
        Paragraph(f"Periode: <b>{period}</b> {sep} Dicetak oleh: {printed_by} {sep} {now} WIB", st["small"]),
        Paragraph("Stok &amp; Penggunaan per Gudang", st["h2"]),
        _pdf_warehouse_table(r, st),
        Spacer(1, 6),
        Paragraph("Kiriman Supplier", st["h2"]),
        _pdf_supplier_table(r, st),
    ]
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=12 * mm, rightMargin=12 * mm, topMargin=12 * mm, bottomMargin=12 * mm,
                            title="Laporan Stok Gudang - TWMS")
    doc.build(story, onFirstPage=_pdf_footer, onLaterPages=_pdf_footer)
    buf.seek(0)
    return buf


@router.get("/export.pdf")
async def export_pdf(date_from: Optional[str] = None, date_to: Optional[str] = None, user: dict = Depends(require_any("stok"))):
    r = await build_report(date_from, date_to, allowed_warehouses(user))
    buf = build_report_pdf(r, period_label(date_from, date_to), user["name"])
    return StreamingResponse(
        buf,
        media_type="application/pdf",
        headers={"Content-Disposition": f'inline; filename="{_filename("pdf", date_from, date_to)}"'},
    )
